from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path("/Users/melvinjose/Desktop/WeatherForecast")
ASSETS = ROOT / "report_assets"
OUT_DOCX = ROOT / "WeatherForecast_System_Report_Melvin_Jose.docx"
OUT_HTML = ROOT / "WeatherForecast_System_Report_Melvin_Jose.html"

figures = [
    (
        ASSETS / "fig1_start.png",
        "Figure 1: Input / Interface / Start Screen (Home page with search and real-time insights panel).",
    ),
    (
        ASSETS / "fig2_processing.png",
        "Figure 2: Core Processing / Intermediate Output (Insight module and temperature panel after location detection).",
    ),
    (
        ASSETS / "fig3_final.png",
        "Figure 3: Final Output / Result (WeatherForecast system output after feature execution).",
    ),
]


cover_lines = [
    "MINI PROJECT REPORT",
    "WeatherForecast System",
    "",
    "Student Name: Melvin Jose",
    "Register Number: AJC140CS23",
    "Class: CSE-C",
    "Group No: Individual",
    "Department: Computer Science and Engineering",
]

abstract = (
    "The WeatherForecast System is a Django-based web application developed to provide accurate, "
    "user-friendly weather updates through both manual city/ZIP search and geolocation-based insights. "
    "The system integrates real-time weather APIs, caches responses to reduce latency and API usage, and "
    "presents data using a responsive glassmorphism interface. In addition to current conditions and "
    "multi-day forecast cards, it includes actionable features such as rain-start prediction, outdoor "
    "planning guidance, air-quality tips, and city-to-city comparison. The project emphasizes clean "
    "architecture by separating service logic from views and enhancing usability with autocomplete, "
    "instant unit switching, theme toggle, and robust error handling for invalid locations and network issues."
)

problem_description = [
    "Weather information is often distributed across multiple apps and websites, making it difficult for users to quickly make daily decisions.",
    "Basic forecast pages may show raw temperature only, without contextual insights such as rain timing, outdoor suitability, and air-quality advisories.",
    "Frequent API calls can also cause slow performance and rate-limit issues if caching is not implemented.",
    "This project solves these issues by providing a single, insight-driven interface that combines real-time data, location awareness, and optimized caching.",
]

libraries_tools = [
    ("Django (5.x)", "Web framework used for URL routing, view rendering, template handling, and project structure management."),
    ("requests", "Used in service layer to fetch weather, forecast, geocoding, and air-pollution data from external APIs."),
    ("python-dotenv", "Loads API key and configuration values from .env for secure secret management."),
    ("OpenWeatherMap APIs", "Primary data provider for current weather, forecast, geolocation lookup, and AQI inputs."),
    ("Django Cache (LocMemCache)", "Stores API responses temporarily to reduce repeated calls and improve speed."),
    ("Bootstrap + Custom CSS", "Provides responsive layout foundation and custom visual design."),
    ("JavaScript (Vanilla)", "Implements geolocation flow, autocomplete, theme toggling, and dynamic insight interactions."),
]

workflow = [
    "1. User opens the homepage and enters a location or uses location permission.",
    "2. Frontend sends query or coordinates to Django endpoint.",
    "3. View delegates weather retrieval to services.py.",
    "4. Service checks cache; if not available, calls OpenWeather APIs.",
    "5. Service parses required fields and computes insight summaries.",
    "6. Django renders structured context into templates.",
    "7. UI displays current temperature, conditions, forecast, and feature outputs.",
]

implementation_details = [
    "Project follows modular architecture with weather_project (settings/routing) and forecast app (views/services/templates).",
    "Services layer isolates API integration, response parsing, temperature conversion, and insight generation logic.",
    "Search endpoint supports both text query and coordinate-based fetch mode.",
    "Autocomplete endpoint returns suggestions for user-typed location strings.",
    "Feature insights endpoint returns domain-specific outputs (rain alert, outdoor planner, AQ tips, city compare).",
    "Current-summary endpoint supports homepage temperature card and fast location-based rendering.",
    "Frontend enforces manual-search validation (\"Type place\") and separates it from \"Use my location\" flow.",
    "Theme system supports standard mode and storm-cloud blue mode with persistent preference.",
    "Client-side and server-side caching are both used to improve perceived performance.",
]

results_text = [
    "The system successfully displays real-time weather information with user-friendly visualization and responsive behavior.",
    "Manual search and geolocation both work with validation and fallback handling.",
    "Insight panels generate contextual outputs that are more actionable than raw weather metrics.",
    "Caching improves repeated-load speed and reduces redundant API requests.",
]

conclusion_scope = [
    "The WeatherForecast System demonstrates a complete end-to-end mini project covering frontend UI, backend services, API integration, caching, and user-centric features.",
    "The current version meets core requirements and provides practical weather insights with good usability.",
    "Scope for extension in mini project phase includes: push notifications for weather alerts, user accounts with favorite cities, map-based radar visualization, multilingual interface, and historical weather analytics dashboards.",
]

references = [
    "Django Documentation: https://docs.djangoproject.com/",
    "OpenWeatherMap API Documentation: https://openweathermap.org/api",
    "Requests Documentation: https://requests.readthedocs.io/",
    "python-dotenv Documentation: https://pypi.org/project/python-dotenv/",
    "Bootstrap Documentation: https://getbootstrap.com/docs/",
]


def add_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def add_bullets(doc: Document, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_docx():
    doc = Document()
    normal = doc.styles["Normal"].font
    normal.name = "Calibri"
    normal.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("MINI PROJECT REPORT")
    run.bold = True
    run.font.size = Pt(20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("WeatherForecast System")
    run.bold = True
    run.font.size = Pt(22)

    doc.add_paragraph("")
    for line in cover_lines[3:]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    add_heading(doc, "2. Abstract")
    doc.add_paragraph(abstract)

    add_heading(doc, "3. Problem Description")
    add_bullets(doc, problem_description)

    add_heading(doc, "4. Libraries & Tools Used (with purpose)")
    add_bullets(doc, [f"{name}: {purpose}" for name, purpose in libraries_tools])

    add_heading(doc, "5. System Design / Workflow")
    add_bullets(doc, workflow)

    add_heading(doc, "6. Implementation Details")
    add_bullets(doc, implementation_details)

    add_heading(doc, "7. Results & Output")
    add_bullets(doc, results_text)

    for image_path, caption in figures:
        if image_path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(image_path), width=Inches(6.2))
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True

    add_heading(doc, "8. Conclusion & Scope for Mini Project Extension")
    add_bullets(doc, conclusion_scope)

    add_heading(doc, "9. References")
    for idx, ref in enumerate(references, start=1):
        doc.add_paragraph(f"{idx}. {ref}")

    doc.save(OUT_DOCX)


def build_html():
    html_parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        "<title>WeatherForecast System Report</title>",
        "<style>",
        "body{font-family:Calibri,Arial,sans-serif;line-height:1.5;margin:36px;color:#111;}",
        "h1,h2,h3{margin-top:22px;margin-bottom:10px;}",
        "h1{font-size:28px;text-align:center;}",
        "h2{font-size:20px;border-bottom:1px solid #ddd;padding-bottom:4px;}",
        "ul{margin:6px 0 8px 18px;}",
        "li{margin-bottom:5px;}",
        ".cover{margin-top:70px;text-align:center;}",
        ".cover p{margin:7px 0;font-size:16px;}",
        ".figure{margin:18px 0;text-align:center;}",
        ".figure img{max-width:100%;height:auto;border:1px solid #bbb;}",
        ".caption{font-size:13px;font-style:italic;margin-top:6px;}",
        ".pagebreak{page-break-before:always;}",
        "</style></head><body>",
    ]
    html_parts.append("<div class='cover'>")
    html_parts.append("<h1>MINI PROJECT REPORT</h1>")
    html_parts.append("<h1>WeatherForecast System</h1>")
    for line in cover_lines[3:]:
        label, value = line.split(":", 1)
        html_parts.append(f"<p><strong>{label}:</strong> {value.strip()}</p>")
    html_parts.append("</div>")
    html_parts.append("<div class='pagebreak'></div>")
    html_parts.append("<h2>2. Abstract</h2>")
    html_parts.append(f"<p>{abstract}</p>")

    html_parts.append("<h2>3. Problem Description</h2><ul>")
    for item in problem_description:
        html_parts.append(f"<li>{item}</li>")
    html_parts.append("</ul>")

    html_parts.append("<h2>4. Libraries & Tools Used (with purpose)</h2><ul>")
    for name, purpose in libraries_tools:
        html_parts.append(f"<li><strong>{name}:</strong> {purpose}</li>")
    html_parts.append("</ul>")

    html_parts.append("<h2>5. System Design / Workflow</h2><ul>")
    for item in workflow:
        html_parts.append(f"<li>{item}</li>")
    html_parts.append("</ul>")

    html_parts.append("<h2>6. Implementation Details</h2><ul>")
    for item in implementation_details:
        html_parts.append(f"<li>{item}</li>")
    html_parts.append("</ul>")

    html_parts.append("<h2>7. Results & Output</h2><ul>")
    for item in results_text:
        html_parts.append(f"<li>{item}</li>")
    html_parts.append("</ul>")

    for path, caption in figures:
        rel = path.relative_to(ROOT)
        html_parts.append("<div class='figure'>")
        html_parts.append(f"<img src='{rel.as_posix()}' alt='{caption}'>")
        html_parts.append(f"<div class='caption'>{caption}</div>")
        html_parts.append("</div>")

    html_parts.append("<h2>8. Conclusion & Scope for Mini Project Extension</h2><ul>")
    for item in conclusion_scope:
        html_parts.append(f"<li>{item}</li>")
    html_parts.append("</ul>")

    html_parts.append("<h2>9. References</h2><ol>")
    for ref in references:
        html_parts.append(f"<li>{ref}</li>")
    html_parts.append("</ol>")
    html_parts.append("</body></html>")

    OUT_HTML.write_text("\n".join(html_parts), encoding="utf-8")


def main():
    build_docx()
    build_html()
    print(f"Generated DOCX: {OUT_DOCX}")
    print(f"Generated HTML: {OUT_HTML}")


if __name__ == "__main__":
    main()
